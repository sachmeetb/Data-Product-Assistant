"""
create_multi_source_samples.py — Generates multi-source banking requirements workbooks for Bronze Agent.
Sources:
  1. Hogan Deposit System (Mainframe Copybooks DD-REC & CI-REC)
  2. SWIFT MT/MX & UK Faster Payments (FPS)
  3. SAP S/4HANA FI-CO General Ledger (BKPF & BSEG)
  4. Salesforce CRM & Client KYC Ingestion
"""

import os
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
SAMPLES_DIR = ROOT / "samples"
SAMPLES_DIR.mkdir(exist_ok=True)

# Styling palette: Bronze Theme
header_fill = PatternFill(start_color="8B4513", end_color="8B4513", fill_type="solid") # SaddleBrown
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
title_fill = PatternFill(start_color="5C2E0B", end_color="5C2E0B", fill_type="solid")  # Dark Bronze
title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
zebra_fill = PatternFill(start_color="FDF5E6", end_color="FDF5E6", fill_type="solid")  # OldLace
thin_side = Side(style="thin", color="D2B48C")
border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)


def style_sheet(ws, title, headers, rows):
    ws.merge_cells("A1:E1")
    ws["A1"] = title
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    ws.row_dimensions[3].height = 24
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for r_idx, row in enumerate(rows, 4):
        ws.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = border
            if is_even:
                cell.fill = zebra_fill

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


# ==============================================================================
# 1. Hogan Deposit System Workbook
# ==============================================================================
wb_hogan = openpyxl.Workbook()

# Sheet 1: Overview
ws1 = wb_hogan.active
ws1.title = "Bank & Ingestion Scope"
hogan_overview = [
    ("Bank Name", "Apex International Bank", "Global commercial and retail banking institution"),
    ("Bank Code", "AIB", "Used for table names and partition prefixes"),
    ("Operating Regions", "UK, EU, US", "Subject to PRA, FCA, GDPR, and Basel III"),
    ("Banking Type", "Retail and Commercial Banking", "Dual domain covering deposit accounts & lending"),
    ("Core Ingestion Source", "Hogan Deposit System", "IBM z/OS Mainframe core banking engine"),
    ("Data Feeds", "DD-REC (Demand Deposits), CI-REC (Customer Information)", "Fixed-width EBCDIC/ASCII records with COBOL copybooks"),
    ("Target Lakehouse", "Google Cloud Data Lake (GCS)", "BigLake Apache Iceberg external tables"),
    ("Retention Policy", "7 Years (2555 Days)", "Mandatory regulatory retention under Basel III / PRA SS34/15"),
    ("Data Standards", "COBOL Copybooks, BIAN v12, ISO 20022", "Zero-loss source fidelity landing"),
]
style_sheet(ws1, "BRONZE INGESTION REQUIREMENT — HOGAN DEPOSIT SYSTEM", ["Property", "Specification Value", "Engineering Notes"], hogan_overview)

# Sheet 2: Feeds & Copybook Fields
ws2 = wb_hogan.create_sheet(title="Hogan Feeds & Copybooks")
hogan_fields = [
    ("brz_hogan_dd_account", "DD-REC", "DDBASIC.CPY", "DD_ACCT_NUM", "STRING", "PIC X(10)", "Primary Account Number"),
    ("brz_hogan_dd_account", "DD-REC", "DDBASIC.CPY", "DD_CUST_NUM", "STRING", "PIC X(12)", "CIS Customer Number Reference"),
    ("brz_hogan_dd_account", "DD-REC", "DDBASIC.CPY", "DD_PROD_TYPE", "STRING", "PIC X(4)", "Product Code (CHKG, SAVG, MMKT)"),
    ("brz_hogan_dd_account", "DD-REC", "DDBASIC.CPY", "DD_STATUS_CD", "STRING", "PIC X(2)", "Account Status (01=Active, 05=Dormant)"),
    ("brz_hogan_dd_account", "DD-REC", "DDBASIC.CPY", "DD_CURR_BAL", "STRING", "PIC S9(11)V99 COMP-3", "Current Ledger Balance"),
    ("brz_hogan_dd_account", "DD-REC", "DDBASIC.CPY", "DD_AVAIL_BAL", "STRING", "PIC S9(11)V99 COMP-3", "Available Balance"),
    ("brz_hogan_ci_customer", "CI-REC", "CIBASIC.CPY", "CI_CUST_NUM", "STRING", "PIC X(12)", "Unique CIS Master Identifier"),
    ("brz_hogan_ci_customer", "CI-REC", "CIBASIC.CPY", "CI_LEGAL_NAME", "STRING", "PIC X(60)", "Full Legal Name (PII)"),
    ("brz_hogan_ci_customer", "CI-REC", "CIBASIC.CPY", "CI_TAX_ID", "STRING", "PIC X(15)", "National Tax Identification (PII)"),
    ("brz_hogan_ci_customer", "CI-REC", "CIBASIC.CPY", "CI_BIRTH_DT", "STRING", "PIC 9(8)", "Date of Birth (YYYYMMDD)"),
]
style_sheet(ws2, "HOGAN MAINFRAME COPYBOOK FIELD MAPPINGS", ["Target Table", "Feed Name", "Copybook", "Field Name", "Type", "PIC Clause", "Description"], hogan_fields)
wb_hogan.save(SAMPLES_DIR / "sample_hogan_deposit_requirements.xlsx")


# ==============================================================================
# 2. SWIFT MT/MX & Faster Payments Workbook
# ==============================================================================
wb_swift = openpyxl.Workbook()
ws_s1 = wb_swift.active
ws_s1.title = "Payments Scope"
swift_overview = [
    ("Bank Name", "Apex International Bank", "Global commercial and retail banking institution"),
    ("Bank Code", "AIB", "Used for table names and partition prefixes"),
    ("Operating Regions", "UK, EU", "Subject to PSR, PRA, FCA, and ISO 20022 mandates"),
    ("Banking Type", "Commercial & Corporate Payments", "Cross-border clearing & domestic real-time payments"),
    ("Core Payment Rails", "SWIFT MT FIN, SWIFT MX (ISO 20022), Faster Payments (FPS)", "Dual-rail payment execution"),
    ("Ingestion Pattern", "Real-time Streaming (Kafka / GCS PubSub)", "Hourly & real-time landing"),
    ("Target Lakehouse", "Google Cloud Data Lake / BigLake Iceberg", "Enforcing 7-year regulatory retention"),
    ("Data Standards", "ISO 20022 pacs.008, SWIFT MT103, BIAN", "Preserving raw FIN and XML payloads"),
]
style_sheet(ws_s1, "BRONZE INGESTION REQUIREMENT — SWIFT & FASTER PAYMENTS", ["Property", "Specification Value", "Engineering Notes"], swift_overview)
wb_swift.save(SAMPLES_DIR / "sample_swift_payments_requirements.xlsx")


# ==============================================================================
# 3. SAP S/4HANA FI-CO General Ledger Workbook
# ==============================================================================
wb_sap = openpyxl.Workbook()
ws_sap1 = wb_sap.active
ws_sap1.title = "SAP ERP Scope"
sap_overview = [
    ("Bank Name", "Apex International Bank", "Global commercial and retail banking institution"),
    ("Bank Code", "AIB", "Used for table names and partition prefixes"),
    ("Operating Regions", "Global", "Subject to SOX, IFRS 9, and Basel III"),
    ("Core Ingestion Source", "SAP S/4HANA (FI-CO General Ledger)", "Enterprise ERP financial ledger"),
    ("Replication Mechanism", "SAP SLT Change Data Capture & IDoc", "Continuous CDC replication to Bronze"),
    ("Target Entities", "BKPF (Document Header), BSEG (Document Line Items)", "Raw financial posting entries"),
    ("Target Lakehouse", "BigLake Apache Iceberg Tables", "7-Year regulatory retention"),
]
style_sheet(ws_sap1, "BRONZE INGESTION REQUIREMENT — SAP GENERAL LEDGER", ["Property", "Specification Value", "Engineering Notes"], sap_overview)
wb_sap.save(SAMPLES_DIR / "sample_sap_general_ledger_requirements.xlsx")


# ==============================================================================
# 4. Salesforce CRM & KYC Ingestion Workbook
# ==============================================================================
wb_sfdc = openpyxl.Workbook()
ws_sfdc1 = wb_sfdc.active
ws_sfdc1.title = "Salesforce Scope"
sfdc_overview = [
    ("Bank Name", "Apex International Bank", "Global commercial and retail banking institution"),
    ("Bank Code", "AIB", "Used for table names and partition prefixes"),
    ("Core Ingestion Source", "Salesforce CRM & Onboarding", "Customer interaction and compliance records"),
    ("Replication Mechanism", "Salesforce Bulk API 2.0 & Pub/Sub Change Data Capture", "Daily batch & real-time CDC"),
    ("Target Entities", "Account, KYC_Verification__c", "Customer organization and compliance signoffs"),
    ("Compliance Controls", "GDPR Article 6, PII Masking Flags, Dataplex Tagging", "Automated sensitive field detection"),
    ("Target Lakehouse", "Google Data Lake / BigLake Iceberg", "5-Year customer interaction retention"),
]
style_sheet(ws_sfdc1, "BRONZE INGESTION REQUIREMENT — SALESFORCE CRM & KYC", ["Property", "Specification Value", "Engineering Notes"], sfdc_overview)
wb_sfdc.save(SAMPLES_DIR / "sample_salesforce_crm_requirements.xlsx")

print("Successfully generated all 4 multi-source sample workbooks in samples/ directory!")
