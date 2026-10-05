#!/usr/bin/env python3
"""
Generate architecture.pptx — 3 slides from the three arch excalidraw diagrams.

Reads each .excalidraw JSON, scales all shapes/arrows/text proportionally
to fill a widescreen slide (13.333" × 7.5"), and assembles into one deck.

Usage:
    python generate_architecture_pptx.py
"""

import sys
import os
import json
import re

sys.path.insert(0, os.path.expanduser("~/.claude/skills/excalidraw-to-pptx/scripts"))

from pptx_lib import (
    add_box, add_text, add_arrow,
    new_presentation, verify_and_save,
    rgb, hex_bare, C,
)
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN

# ─── Slide layout constants ───────────────────────────────────────────────────

SLIDE_W   = 13.333   # inches
SLIDE_H   = 7.5
TITLE_H   = 0.95     # reserved for slide title + rule
MARGIN    = 0.3
CONTENT_L = MARGIN
CONTENT_T = TITLE_H
CONTENT_W = SLIDE_W - 2 * MARGIN
CONTENT_H = SLIDE_H - TITLE_H - MARGIN

# ─── Color helpers ────────────────────────────────────────────────────────────

def _hex(color, fallback="CBD5E1"):
    """Normalize any color string to a 6-digit hex (no #). Returns None for transparent."""
    if not color or color in ("transparent", "none", ""):
        return None
    c = str(color).strip()
    if c.startswith("#"):
        h = c[1:]
        if len(h) == 3:
            return "".join(x * 2 for x in h)
        return h[:6] if len(h) >= 6 else fallback
    m = re.match(r"rgba?\(\s*(\d+),\s*(\d+),\s*(\d+)", c)
    if m:
        return "{:02x}{:02x}{:02x}".format(int(m[1]), int(m[2]), int(m[3]))
    named = {"white": "FFFFFF", "black": "000000", "red": "FF0000",
             "green": "00FF00", "blue": "0000FF"}
    return named.get(c.lower(), fallback)


# ─── Bounds and scaling ───────────────────────────────────────────────────────

def _bounds(elements):
    """Compute bounding box of all visible elements."""
    xs, ys = [], []
    for el in elements:
        if el.get("isDeleted"):
            continue
        t = el.get("type", "")
        if t in ("rectangle", "ellipse", "diamond"):
            x, y = el.get("x", 0), el.get("y", 0)
            w, h = el.get("width", 0), el.get("height", 0)
            xs += [x, x + w]
            ys += [y, y + h]
        elif t == "text":
            xs.append(el.get("x", 0))
            ys.append(el.get("y", 0))
        elif t in ("arrow", "line"):
            bx, by = el.get("x", 0), el.get("y", 0)
            for p in el.get("points", []):
                xs.append(bx + p[0])
                ys.append(by + p[1])
    if not xs:
        return 0, 0, 1440, 900
    return min(xs), min(ys), max(xs), max(ys)


def _make_scaler(elements):
    """Return coord-conversion functions (ix, iy, iw) for scaling to slide content area."""
    x1, y1, x2, y2 = _bounds(elements)
    cw = max(x2 - x1, 1)
    ch = max(y2 - y1, 1)
    sx = CONTENT_W / cw
    sy = CONTENT_H / ch
    scale = min(sx, sy)

    actual_w = cw * scale
    actual_h = ch * scale
    off_x = CONTENT_L + (CONTENT_W - actual_w) / 2
    off_y = CONTENT_T + (CONTENT_H - actual_h) / 2

    def ix(px):
        return Inches(off_x + (px - x1) * scale)

    def iy(py):
        return Inches(off_y + (py - y1) * scale)

    def iw(px):
        return Inches(max(abs(px) * scale, 0.02))

    def fs(px):
        """Scale SVG font px to PPTX pt, clamped to readable range."""
        return max(6, min(16, round(abs(px) * scale * 72)))

    return ix, iy, iw, fs


# ─── Slide builder ────────────────────────────────────────────────────────────

def _add_slide_title(slide, title, subtitle):
    """Add slide title + subtitle + rule to the top of the slide."""
    # Title
    add_text(slide,
             left=Inches(CONTENT_L), top=Inches(0.1),
             width=Inches(CONTENT_W), height=Inches(0.45),
             text_str=title,
             font_size=22, font_color=C["dark"], bold=True,
             align=PP_ALIGN.LEFT)
    # Subtitle
    add_text(slide,
             left=Inches(CONTENT_L), top=Inches(0.56),
             width=Inches(CONTENT_W), height=Inches(0.22),
             text_str=subtitle,
             font_size=9, font_color=C["body_text"],
             align=PP_ALIGN.LEFT)
    # Rule line
    add_arrow(slide,
              [(Inches(CONTENT_L), Inches(TITLE_H - 0.08)),
               (Inches(CONTENT_L + CONTENT_W), Inches(TITLE_H - 0.08))],
              color_hex=C["skills_border"],
              width_pt=1.2, arrowhead=False)


def build_slide(slide, excalidraw_path, title, subtitle):
    """Map all elements from an Excalidraw file onto a PPTX slide."""
    with open(excalidraw_path) as f:
        doc = json.load(f)

    elements = [e for e in doc.get("elements", []) if not e.get("isDeleted")]
    ix, iy, iw, fs = _make_scaler(elements)

    _add_slide_title(slide, title, subtitle)

    # Render in element order (SVG painter model: earlier = behind)
    for el in elements:
        t = el.get("type", "")

        if t == "rectangle":
            fill = _hex(el.get("backgroundColor", "#ffffff"))
            if fill is None:
                continue                      # skip transparent-fill rects
            stroke = _hex(el.get("strokeColor", "#ced4da")) or "CBD5E1"
            sw_px = el.get("strokeWidth", 1)
            sw_pt = max(0.3, sw_px * 0.6)    # px → pt (approximate visual weight)
            try:
                add_box(slide,
                        left=ix(el["x"]), top=iy(el["y"]),
                        width=iw(el.get("width", 1)),
                        height=iw(el.get("height", 1)),
                        fill_hex=fill, border_hex=stroke,
                        text_str="",
                        border_width=Pt(sw_pt))
            except Exception:
                pass

        elif t == "text":
            txt = el.get("text", "").strip()
            if not txt:
                continue
            color = _hex(el.get("strokeColor", "#1e1e1e")) or "1e1e1e"
            align_str = el.get("textAlign", "left")
            align = (PP_ALIGN.CENTER if align_str == "center"
                     else PP_ALIGN.RIGHT if align_str == "right"
                     else PP_ALIGN.LEFT)
            try:
                add_text(slide,
                         left=ix(el["x"]), top=iy(el["y"]),
                         width=iw(max(el.get("width", 10), 10)),
                         height=iw(max(el.get("height", 10), 10)),
                         text_str=txt,
                         font_size=fs(el.get("fontSize", 11)),
                         font_color="#" + color,
                         align=align)
            except Exception:
                pass

        elif t in ("arrow", "line"):
            bx, by = el.get("x", 0), el.get("y", 0)
            points = el.get("points", [])
            if len(points) < 2:
                continue
            pptx_pts = [(ix(bx + p[0]), iy(by + p[1])) for p in points]
            stroke = _hex(el.get("strokeColor", "#868e96")) or "868e96"
            sw_pt = max(0.5, el.get("strokeWidth", 1.5) * 0.6)
            dashed = el.get("strokeStyle", "solid") == "dashed"
            has_arrow = el.get("endArrowhead") == "arrow"
            try:
                add_arrow(slide, pptx_pts,
                          color_hex="#" + stroke,
                          width_pt=sw_pt,
                          dashed=dashed, arrowhead=has_arrow)
            except Exception:
                pass


# ─── Main ─────────────────────────────────────────────────────────────────────

def build():
    HERE = os.path.dirname(os.path.abspath(__file__))
    OUT  = os.path.join(HERE, "architecture.pptx")

    prs, slide1 = new_presentation()

    build_slide(slide1,
                os.path.join(HERE, "arch-v1-full-stack.excalidraw"),
                "Data Workbench — System Architecture",
                "Full layered stack · UI shells → Agent Harness → Skills Ecosystem → Data & Memory")

    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.background.fill.solid()
    slide2.background.fill.fore_color.rgb = rgb(C["white"])
    build_slide(slide2,
                os.path.join(HERE, "arch-v2-skills-ecosystem.excalidraw"),
                "Agent Skills Ecosystem",
                "Skill anatomy, catalog depth, and harness orchestration model")

    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    slide3.background.fill.solid()
    slide3.background.fill.fore_color.rgb = rgb(C["white"])
    build_slide(slide3,
                os.path.join(HERE, "arch-v3-data-product-journey.excalidraw"),
                "Data Product Journey",
                "Six-phase lifecycle · PO ideation → Engineering → Marketplace deployment")

    verify_and_save(prs, slide1, OUT)


if __name__ == "__main__":
    build()
