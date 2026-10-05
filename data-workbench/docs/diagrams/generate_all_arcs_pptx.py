#!/usr/bin/env python3
"""
generate_all_arcs_pptx.py
Converts all arc SVG files → single PPTX with native python-pptx shapes.
No colors: transparent fills, gray borders, black text.
"""
import os
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SVG_W, SVG_H   = 1440.0, 810.0
SLIDE_W, SLIDE_H = 12192000, 6858000   # 16:9 widescreen EMU
SX = SLIDE_W / SVG_W    # 8466.67 EMU/px
SY = SLIDE_H / SVG_H

GRAY = RGBColor(0xBB, 0xBB, 0xBB)
INK  = RGBColor(0x1E, 0x29, 0x3B)


def flt(v, d=0.0):
    try:
        return float(str(v).replace("px", "").strip())
    except Exception:
        return float(d)


def local(el):
    t = el.tag
    return t.split("}")[-1] if "}" in t else t


def collect_text(el):
    parts = []
    if el.text and el.text.strip():
        parts.append(el.text.strip())
    for child in el:
        if child.text and child.text.strip():
            parts.append(child.text.strip())
    return " ".join(parts)


def add_slide(prs, svg_path, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank

    # Small slide-title label top-right
    lbl = slide.shapes.add_textbox(
        Emu(int(SLIDE_W * 0.65)), Emu(0),
        Emu(int(SLIDE_W * 0.34)), Emu(int(18 * SY))
    )
    r = lbl.text_frame.paragraphs[0].add_run()
    r.text = title
    r.font.size = Pt(6)
    r.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    lbl.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    tree = ET.parse(svg_path)
    root = tree.getroot()
    n_shapes = 0

    for el in root.iter():
        tag = local(el)

        # ── RECT ──────────────────────────────────────────────────────────────
        if tag == "rect":
            w = flt(el.get("width", 0))
            h = flt(el.get("height", 0))
            if w < 2 or h < 2:
                continue
            x = flt(el.get("x", 0))
            y = flt(el.get("y", 0))

            shp = slide.shapes.add_shape(
                1,                              # 1 = RECTANGLE in MSO_AUTO_SHAPE_TYPE
                Emu(int(x * SX)),
                Emu(int(y * SY)),
                Emu(max(int(w * SX), 5000)),
                Emu(max(int(h * SY), 5000)),
            )
            shp.fill.background()
            shp.line.color.rgb = GRAY
            shp.line.width = Pt(0.5)
            n_shapes += 1

        # ── TEXT ──────────────────────────────────────────────────────────────
        elif tag == "text":
            txt = collect_text(el)
            if not txt:
                continue

            xp     = flt(el.get("x", 0))
            yp     = flt(el.get("y", 0))
            fs_px  = flt(el.get("font-size", 10))
            fw     = el.get("font-weight", "normal")
            anchor = el.get("text-anchor", "start")

            # Estimate box size
            fs_emu    = int(fs_px * SY)
            char_w    = int(fs_emu * 0.60)
            est_w     = max(len(txt) * char_w, 180000)
            est_h     = max(int(fs_emu * 1.8), 50000)

            # Anchor → left edge
            lx = int(xp * SX)
            if anchor == "middle":
                lx -= est_w // 2
            elif anchor == "end":
                lx -= est_w
            lx = max(0, lx)

            # SVG y is baseline → shift up by ~one line
            ty = max(0, int(yp * SY) - fs_emu)

            tb = slide.shapes.add_textbox(
                Emu(lx), Emu(ty), Emu(est_w), Emu(est_h)
            )
            tf   = tb.text_frame
            tf.word_wrap = False
            para = tf.paragraphs[0]
            run  = para.add_run()
            run.text = txt

            fs_pt          = max(6, round(fs_px * 0.70))
            run.font.size  = Pt(fs_pt)
            run.font.bold  = fw in ("700", "600", "bold")
            run.font.color.rgb = INK

            if anchor == "middle":
                para.alignment = PP_ALIGN.CENTER
            elif anchor == "end":
                para.alignment = PP_ALIGN.RIGHT

            n_shapes += 1

    print(f"  {title}: {n_shapes} shapes added from {svg_path}")
    return slide


def main():
    arc_files = [
        ("arc_intro_lead.svg",           "Intro — What is Data Workbench"),
        ("arc_features_detail.svg",       "How It Works — Agent Lifecycle"),
        ("arc1_ai_readiness_gap.svg",     "Arc 1 v1 — AI Readiness Gap"),
        ("arc1_ai_readiness_gap_v2.svg",  "Arc 1 v2 — AI Builds the Foundation"),
        ("arc2_months_to_days.svg",       "Arc 2 v1 — DB vs Data Product Pipeline"),
        ("arc2_months_to_days_v2.svg",    "Arc 2 v2 — AI Doing the Data Work"),
        ("arc3_data_product_stack.svg",   "Arc 3 v1 — Three-Layer Stack"),
        ("arc3_data_product_stack_v2.svg","Arc 3 v2 — AI Builder + AI Consumer"),
    ]

    prs = Presentation()
    prs.slide_width  = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)

    found = 0
    for filename, title in arc_files:
        if os.path.exists(filename):
            add_slide(prs, filename, title)
            found += 1
        else:
            print(f"  SKIP (not found): {filename}")

    out = "data_workbench_arcs.pptx"
    prs.save(out)
    print(f"\n✓  {found} slides → {out}")


main()
